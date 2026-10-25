"""Submit render jobs to the network-isolated LibreOffice worker over a spool."""

import json
import os
import tempfile
import time
from io import BytesIO
from pathlib import Path


def sample_pages(total, maximum=8):
    """Source page numbers, including both ends, without relabeling the sample."""
    count = min(total, maximum)
    if count < 1:
        return ()
    if count == 1:
        return (1,)
    return tuple(1 + i * (total - 1) // (count - 1) for i in range(count))


def render_input(data, pages=None):
    """Render an explicit page order, including hidden pages, on a private copy."""
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    numbers = (
        tuple(pages) if pages is not None else tuple(range(1, len(prs.slides) + 1))
    )
    if (
        not numbers
        or len(set(numbers)) != len(numbers)
        or any(not isinstance(n, int) or not 1 <= n <= len(prs.slides) for n in numbers)
    ):
        raise ValueError("渲染页码无效")
    slide_ids = prs.slides._sldIdLst
    chosen = [slide_ids[n - 1] for n in numbers]
    for node in list(slide_ids):
        slide_ids.remove(node)
    for node in chosen:
        slide_ids.append(node)
    for slide in prs.slides:
        slide._element.set("show", "1")
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def render_pdf(pptx_bytes):
    spool = os.environ.get("LANDPPT_PPTX_RENDER_SPOOL", "")
    if not spool or not Path(spool).is_dir():
        raise ValueError(
            "PPTX 渲染服务未配置，请按模板包 README 启动隔离渲染容器并设置 LANDPPT_PPTX_RENDER_SPOOL"
        )
    # Random job directories contain no user-chosen paths. The worker writes only here.
    with tempfile.TemporaryDirectory(prefix="pptx-", dir=spool) as folder:
        job = Path(folder)
        (job / "input.pptx").write_bytes(pptx_bytes)
        (job / "ready").touch()
        deadline = time.monotonic() + 240
        while time.monotonic() < deadline:
            status = job / "result.json"
            if status.exists():
                result = json.loads(status.read_text(encoding="utf-8"))
                if result.get("error"):
                    raise ValueError(result["error"])
                pdf = job / "input.pdf"
                if not pdf.exists() or pdf.stat().st_size > 80_000_000:
                    raise ValueError("渲染结果缺失或过大")
                return pdf.read_bytes(), result.get("fonts", [])
            time.sleep(0.25)
        raise ValueError("PPTX 渲染超时，请检查渲染服务后重试")
