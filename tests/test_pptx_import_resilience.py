"""Object reconstruction, exact source references and per-page failure isolation."""

import base64
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches
from test_pptx_ai_analysis import analysis_page, plan
from test_template_package_pptx_import import fake_render, fixture_pptx, png

from landppt.services.template.global_master_template_service import (
    GlobalMasterTemplateService,
)
from landppt.services.template_package import (
    pptx_analysis,
    pptx_import,
    pptx_render,
    pptx_vision,
)
from landppt.services.template_package.service import validate_package
from landppt.services.slide.package_generation.candidate_filter import display_fields


def saved(prs):
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def test_render_copy_includes_hidden_pages_and_preserves_source_order():
    raw = fixture_pptx(3)
    result = pptx_render.render_input(raw, [3, 2])
    selected = Presentation(BytesIO(result))
    assert [s.shapes[0].text for s in selected.slides] == [
        "Slide 3 title",
        "Slide 2 title",
    ]
    assert all(s._element.get("show") == "1" for s in selected.slides)
    assert Presentation(BytesIO(raw)).slides[1]._element.get("show") == "0"


@pytest.mark.asyncio
async def test_normal_reference_uses_isolated_renderer_with_sample_source_numbers(
    monkeypatch,
):
    raw = fixture_pptx(12)
    seen = []

    def render(data):
        selected = Presentation(BytesIO(data))
        seen.extend(s.shapes[0].text for s in selected.slides)
        assert all(s._element.get("show") == "1" for s in selected.slides)
        return fake_render(data)

    monkeypatch.setattr(pptx_render, "render_pdf", render)
    service = GlobalMasterTemplateService(user_id=1)
    screenshots = await service._render_pptx_slides_to_images(
        raw, page_numbers=[12, 2, 7]
    )
    assert seen == ["Slide 12 title", "Slide 2 title", "Slide 7 title"]
    assert [s["slide_idx"] for s in screenshots] == [12, 2, 7]
    assert all(s["blob"].startswith(b"\x89PNG") for s in screenshots)


@pytest.mark.asyncio
async def test_normal_extraction_samples_evenly_and_does_not_use_photo_as_page(
    monkeypatch,
):
    service = GlobalMasterTemplateService(user_id=1)
    seen = []

    async def render(raw, **kwargs):
        seen.extend(kwargs["page_numbers"])
        return []  # Unavailable page renderer is a structural-only reference.

    monkeypatch.setattr(service, "_render_pptx_slides_to_images", render)
    monkeypatch.setattr(
        service,
        "_upload_pptx_image_to_hosting",
        AsyncMock(return_value="https://images.test/resource.png"),
    )
    result = await service._extract_pptx_template_reference(
        {"filename": "source.pptx", "data": base64.b64encode(fixture_pptx(12)).decode()}
    )
    assert seen == list(pptx_render.sample_pages(12, 8)) == result["sampled_pages"]
    assert seen[0] == 1 and seen[-1] == 12 and len(seen) == 8
    assert "第12页" in result["analysis_summary"]
    assert result["reference_image"] is None
    assert result["uploaded_images"]  # Resource photos still available for HTML.


def native_source():
    prs = Presentation(BytesIO(fixture_pptx(1)))
    slide = prs.slides[0]
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            shape._element.getparent().remove(shape._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(239, 245, 250)
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.6), Inches(2.8)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(200, 220, 240)
    card.line.fill.background()
    etree.SubElement(card._element.spPr, qn("a:effectLst"))
    card._element.xpath("./p:style/a:effectRef")[0].set("idx", "0")
    tree = slide.shapes._spTree
    tree.remove(card._element)
    tree.insert(2, card._element)
    # A source icon, not a full-page raster or a content photograph.
    slide.shapes.add_picture(
        BytesIO(png("#2878b0")), Inches(0.1), Inches(0.1), Inches(0.25), Inches(0.25)
    )
    return saved(prs)


def options_for(raw, monkeypatch, mode="visual"):
    monkeypatch.setattr(pptx_import, "render_pdf", fake_render)
    analysis = pptx_import.analyze_pptx(raw)
    return pptx_import.ImportOptions(
        name="Source objects",
        mode=mode,
        pages=[
            {
                "slide": s["slide"],
                "bindings": {c["id"]: c["suggested"] for c in s["candidates"]},
            }
            for s in analysis["slides"]
        ],
    )


def test_native_shapes_icons_and_slots_pass_without_model_redrawing(monkeypatch):
    raw = native_source()
    options = options_for(raw, monkeypatch)

    def no_model(*_):
        pytest.fail(f"supported native objects must not need a redraw: {_[1]}")

    package, report = pptx_import.import_pptx(raw, options, reconstruct=no_model)
    assert report["methods"] == [{"slide": 1, "method": "native"}]
    assert not report["partial"]
    svg = package.components[0].svg
    assert "#EFF5FA" in svg and "#C8DCF0" in svg
    assert (
        "Body content 1" in display_fields(package.components[0].examples[0]).values()
    )
    root = etree.fromstring(svg.encode())
    assert any(n.get("data-asset") for n in root.iter())  # The original small icon.
    assert all(
        not (n.text or "").strip() for n in root.iter() if n.tag.endswith("}text")
    )
    assert package.components[0].reference_asset not in [
        n.get("data-asset") for n in root.iter()
    ]
    validate_package(package)


def test_invalid_selected_page_does_not_discard_valid_pages(monkeypatch):
    raw = fixture_pptx(2)
    options = options_for(raw, monkeypatch, mode="preserve")
    invalid = options.pages[0].model_copy(update={"bindings": {"missing:0": "title"}})
    options = options.model_copy(update={"pages": (invalid, options.pages[1])})
    package, report = pptx_import.import_pptx(raw, options)
    assert [c.source_slide for c in package.components] == [2]
    assert report["partial"] and report["failures"][0]["slide"] == 1
    assert "未知槽位" in report["failures"][0]["error"]
    validate_package(package)


def test_unsupported_native_effect_uses_visual_model_instead_of_dropping_it(
    monkeypatch,
):
    prs = Presentation(BytesIO(native_source()))
    card = next(
        s
        for s in prs.slides[0].shapes
        if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    )
    effects = card._element.spPr.find(qn("a:effectLst"))
    etree.SubElement(effects, qn("a:outerShdw"), blurRad="40000")
    raw = saved(prs)
    options = options_for(raw, monkeypatch)
    calls = []

    def reconstruct(component, error, reference):
        calls.append(error)
        root = etree.fromstring(component.svg.encode())
        for node in list(root):
            if node.get("data-asset"):
                root.remove(node)
        return pptx_vision.apply_visual_design(
            component, {"svg": etree.tostring(root, encoding="unicode")}
        )

    package, report = pptx_import.import_pptx(raw, options, reconstruct=reconstruct)
    assert len(calls) == 1 and "阴影" in calls[0]
    assert report["methods"] == [{"slide": 1, "method": "visual"}]
    validate_package(package)


@pytest.mark.asyncio
async def test_failed_analysis_page_retains_diagnostics_and_other_successes(
    monkeypatch,
):
    first, second, third = analysis_page(), analysis_page(), analysis_page()
    second["slide"], third["slide"] = 2, 3
    completion = AsyncMock(
        side_effect=[
            (plan(), {}),
            ValueError("invalid plan"),
            ValueError("invalid plan"),
            ValueError("invalid plan"),
            (plan(), {}),
        ]
    )
    monkeypatch.setattr(pptx_analysis.ContentService, "json_completion", completion)
    result = await pptx_analysis.classify_analysis(
        {"slides": [first, second, third]}, None, lambda *_: None
    )
    assert result["partial"] and result["failures"][0]["slide"] == 2
    assert second["review_required"] and second["analysis_method"] == "failed"
    assert first["analysis_method"] == third["analysis_method"] == "ai"
    assert completion.await_count == 5


@pytest.mark.asyncio
async def test_image_only_page_is_reviewable_without_spending_a_model_call(monkeypatch):
    completion = AsyncMock(return_value=(plan(), {}))
    monkeypatch.setattr(pptx_analysis.ContentService, "json_completion", completion)
    empty = {"slide": 1, "candidates": [], "warnings": []}
    second = {**analysis_page(), "slide": 2}
    result = await pptx_analysis.classify_analysis(
        {"slides": [empty, second]}, None, lambda *_: None
    )
    assert completion.await_count == 1
    assert result["partial"] and empty["review_required"]
    assert "可编辑文字" in empty["warnings"][0]


def test_replacement_capacity_is_measured_even_when_short_source_fits():
    assert pptx_vision.required_text_height(
        "短", "blocks[0].body", 100, 18
    ) > pptx_vision.text_height("短", 100, 18)


@pytest.mark.asyncio
async def test_first_successful_screenshot_is_used_when_first_upload_fails(monkeypatch):
    service = GlobalMasterTemplateService(user_id=1)

    async def render(*_, **__):
        return [
            {"slide_idx": n, "blob": png(), "content_type": "image/png"} for n in (1, 2)
        ]

    async def upload(_, __, name):
        return (
            None if name == "slide1_screenshot" else f"https://images.test/{name}.png"
        )

    monkeypatch.setattr(service, "_render_pptx_slides_to_images", render)
    monkeypatch.setattr(service, "_upload_pptx_image_to_hosting", upload)
    result = await service._extract_pptx_template_reference(
        {
            "data": base64.b64encode(fixture_pptx(2)).decode(),
            "filename": "source.pptx",
        }
    )
    assert result["reference_image"]["filename"] == "slide2_screenshot.png"
    assert result["reference_image"]["data"].endswith("/slide2_screenshot.png")
    assert not result["extra_reference_images"]


@pytest.mark.asyncio
async def test_normal_summary_uses_transformed_group_coordinates(monkeypatch):
    prs = Presentation(BytesIO(fixture_pptx(1)))
    group = prs.slides[0].shapes.add_group_shape([prs.slides[0].shapes[1]])
    group.left, group.top = int(prs.slide_width * 0.4), int(prs.slide_height * 0.3)
    group.width, group.height = int(prs.slide_width * 0.5), int(prs.slide_height * 0.2)
    service = GlobalMasterTemplateService(user_id=1)
    monkeypatch.setattr(
        service, "_render_pptx_slides_to_images", AsyncMock(return_value=[])
    )
    monkeypatch.setattr(
        service, "_upload_pptx_image_to_hosting", AsyncMock(return_value=None)
    )
    result = await service._extract_pptx_template_reference(
        {"data": base64.b64encode(saved(prs)).decode()}
    )
    assert "x=40%,y=30%,w=50%,h=20%" in result["analysis_summary"]


@pytest.mark.asyncio
async def test_all_failed_analysis_is_an_error_not_empty_success(monkeypatch):
    completion = AsyncMock(side_effect=ValueError("invalid plan"))
    monkeypatch.setattr(pptx_analysis.ContentService, "json_completion", completion)
    with pytest.raises(ValueError, match="没有页面完成识别"):
        await pptx_analysis.classify_analysis(
            {"slides": [analysis_page()]}, None, lambda *_: None
        )
    assert completion.await_count == 3
