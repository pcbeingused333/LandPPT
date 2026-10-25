"""Fast source replicas retain real artwork and editable, measured content."""

from io import BytesIO

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from test_pptx_import_resilience import native_source, saved
from test_template_package_pptx_import import fake_render, fixture_pptx

from landppt.services.template_package import pptx_import, pptx_replica
from landppt.services.template_package.service import validate_package
from landppt.services.slide.package_generation.candidate_filter import display_fields


def test_one_step_replica_renders_source_once_and_cleared_artwork_once(monkeypatch):
    renders = []

    def render(data):
        renders.append(data)
        return fake_render(data)

    monkeypatch.setattr(pptx_import, "render_pdf", render)
    source = native_source()
    package, report = pptx_import.replicate_pptx(source, "Faithful")
    assert len(renders) == 2
    assert not report["partial"]
    assert report["methods"] == [{"slide": 1, "method": "replica"}]
    validate_package(package)
    component = package.components[0]
    fields = display_fields(component.examples[0])
    assert "Slide 1 title" == fields["title"]
    assert "Body content 1" in fields.values()
    fixed_assets = [
        n.get("data-asset")
        for n in etree.fromstring(component.svg.encode()).iter()
        if n.get("data-asset")
    ]
    assert component.reference_asset not in fixed_assets
    cleaned = Presentation(BytesIO(renders[1]))
    assert all(
        "Body content 1" not in s.text
        for s in cleaned.slides[0].shapes
        if s.has_text_frame
    )
    assert "Body content 1" in next(
        s.text
        for s in Presentation(BytesIO(source)).slides[0].shapes
        if "Body" in getattr(s, "text", "")
    )


def test_hidden_source_pages_are_not_automatically_imported(monkeypatch):
    monkeypatch.setattr(pptx_import, "render_pdf", fake_render)
    prs = Presentation(BytesIO(fixture_pptx(3)))
    prs.slides[2].shapes[0].left += Inches(0.3)
    package, report = pptx_import.replicate_pptx(saved(prs), "Visible pages")
    assert [c.source_slide for c in package.components] == [1, 3]
    assert not report["partial"]


def candidate(key, text, box, *, suggested="body", size=18, kind="text"):
    return dict(
        id=key,
        text=text,
        box=box,
        suggested=suggested,
        kind=kind,
        font_size=size,
        bold=False,
        font="Arial",
        color="#cccccc",
    )


def test_auto_binding_keeps_short_english_labels_editable_and_overlay_artwork_fixed():
    rows = [
        candidate("1:0", "Title", [30, 20, 1000, 70], suggested="title"),
        candidate("2:0", "RUNNING", [30, 200, 180, 40], suggested="fixed"),
        candidate("3:0", "2 / 15", [1100, 660, 80, 40], suggested="fixed"),
        candidate("4", "panel", [20, 190, 500, 300], kind="image"),
        candidate("5", "photo", [650, 200, 500, 300], kind="image"),
    ]
    result = pptx_replica.automatic_page(2, rows)
    assert result.bindings == {
        "1:0": "title",
        "2:0": "heading",
        "3:0": "fixed",
        "4": "fixed",
        "5": "image",
    }


def test_card_badges_and_paragraphs_are_one_point_without_losing_source_copy():
    rows = [
        candidate("1:0", "Title", [30, 20, 1000, 60], suggested="title"),
        candidate("2:0", "Card name", [80, 200, 200, 28]),
        candidate("3:0", "A complete source paragraph", [80, 240, 300, 50]),
        candidate("4:0", "Label", [80, 310, 70, 28]),
        candidate("5:0", "Another label", [180, 310, 140, 28]),
    ]
    pptx_replica.assign_regions(
        rows, [dict(id="card", box=[60, 180, 360, 200], picture=False, visible=True)]
    )
    plan = pptx_replica.automatic_page(2, rows)
    assert len(set(plan.groups.values())) == 1
    merged = pptx_replica.merge_replica_groups(
        [(c, plan.bindings[c["id"]]) for c in rows], plan.groups
    )
    assert len(merged) == 3  # page title, card heading, card copy
    body = next(c for c, r in merged if r == "body")
    assert body["text"] == "A complete source paragraph · Label · Another label"
    assert body["box"][0] >= 60 and body["box"][0] + body["box"][2] <= 420


def test_image_only_page_is_rejected_instead_of_baked_into_template():
    with pytest.raises(ValueError, match="没有可编辑文字"):
        pptx_replica.automatic_page(
            1, [candidate("1", "scan", [0, 0, 1280, 720], kind="image")]
        )


def test_final_gate_still_rejects_narrow_copy_and_replica_repairs_at_most_once(
    monkeypatch,
):
    monkeypatch.setattr(pptx_import, "render_pdf", fake_render)
    raw = native_source()
    analysis = pptx_import.analyze_pptx(raw)
    choices = [
        pptx_replica.automatic_page(s["slide"], s["candidates"])
        for s in analysis["slides"]
    ]
    settings = pptx_import.ImportOptions(name="Bounded", mode="replica", pages=choices)
    attempts = []
    monkeypatch.setattr(
        pptx_import,
        "check_capacity",
        lambda *_: (_ for _ in ()).throw(ValueError("Too narrow")),
    )

    def repair(component, error):
        attempts.append(error)
        return component

    with pytest.raises(ValueError, match="Too narrow"):
        pptx_import.import_pptx(raw, settings, repair=repair)
    assert len(attempts) == 1


def test_cleared_card_overlay_uses_svg_namespace_and_measured_source_color():
    from test_template_package_pptx_import import package_with_background

    package = package_with_background()
    component = package.components[0]
    root = etree.fromstring(component.svg.encode())
    next(n for n in root if n.get("id") == "import-background").set("id", "background")
    component = component.model_copy(
        update={"svg": etree.tostring(root, encoding="unicode")}
    )
    cleaned = pptx_replica.clean_card_interiors(
        component, [(dict(artwork_box=[100, 200, 200, 100]), "body")], package.assets[0]
    )
    overlays = [
        n
        for n in etree.fromstring(cleaned.svg.encode())
        if n.get("fill") == "#ffffff" and n.get("width") == "200"
    ]
    assert len(overlays) == 1 and overlays[0].tag == "{http://www.w3.org/2000/svg}rect"
    validate_package(package.model_copy(update={"components": (cleaned,)}))


def test_merged_badge_does_not_leave_empty_colored_capsule(monkeypatch):
    prs = Presentation(BytesIO(native_source()))
    badge = prs.slides[0].shapes.add_textbox(
        Inches(2), Inches(3.3), Inches(2), Inches(0.4)
    )
    badge.name = "Source badge"
    badge.text = "Tag"
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(240, 120, 40)
    renders = []

    def render(data):
        renders.append(data)
        return fake_render(data)

    monkeypatch.setattr(pptx_import, "render_pdf", render)
    pptx_import.replicate_pptx(saved(prs), "Clean badges")
    result = Presentation(BytesIO(renders[1]))
    cleared = next(s for s in result.slides[0].shapes if s.name == "Source badge")
    assert not cleared.text
    assert "noFill" in cleared._element.spPr.xml
