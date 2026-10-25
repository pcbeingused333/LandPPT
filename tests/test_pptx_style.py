"""Style replicas learn a deck's look and lay out every family on clean artwork."""

from collections import Counter
from io import BytesIO

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from test_template_package_pptx_import import fake_render, png

from landppt.services.slide.package_generation.capacity import slot_limits
from landppt.services.template_package import pptx_import, pptx_style
from landppt.services.template_package.service import validate_package

ORANGE = RGBColor(0xF5, 0x9E, 0x0B)


def _text(slide, text, left, top, width, height, size, color=RGBColor(40, 40, 40)):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    return box


def style_source(pages=5):
    """Cover, card-based content pages and an ending, like generated decks."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    for index in range(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.08), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ORANGE
        bar.name = "edge bar"
        _text(slide, f"Page {index + 1} title", 0.6, 0.4, 9, 0.8, 32, ORANGE)
        if index in (0, pages - 1):
            continue
        for column in range(3):
            left = 0.6 + column * 4.1
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(2), Inches(3.8), Inches(2.4)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(20, 22, 25)
            card.name = f"card {column}"
            _text(slide, f"Card heading {column}", left + 0.2, 2.2, 3.4, 0.5, 16)
            _text(slide, f"Card body text {column} " * 3, left + 0.2, 2.8, 3.4, 1.2, 12)
        slide.shapes.add_picture(
            BytesIO(png("#2878b0")), Inches(4), Inches(5), Inches(2), Inches(1.5)
        ).name = "content photo"
    output = BytesIO()
    prs.save(output)
    return output.getvalue()


def test_style_replica_covers_every_family_with_two_renders(monkeypatch):
    renders = []

    def render(data):
        renders.append(data)
        return fake_render(data)

    monkeypatch.setattr(pptx_import, "render_pdf", render)
    package, report = pptx_style.style_replicate_pptx(style_source(), "Deck style")
    assert len(renders) == 2  # source once, cleaned artwork once
    assert report["failures"] == [] and not report["partial"]
    validate_package(package)
    families = Counter(c.family for c in package.components)
    assert set(families) == {
        "cover", "section", "points", "comparison", "process", "image", "metrics", "summary"
    }
    assert families["process"] == 3 and families["points"] == 2
    assert {m["method"] for m in report["methods"]} == {"style"}
    # Every layout is checked with short and near-capacity copy, not only previews.
    assert all(len(c.examples) == 3 for c in package.components)
    assert report["style"]["source_pages"] == [1, 2]
    assert report["style"]["card"]["radius"] > 0


def test_artwork_keeps_edge_decoration_and_drops_content_objects(monkeypatch):
    captured = []

    def render(data):
        captured.append(data)
        return fake_render(data)

    monkeypatch.setattr(pptx_import, "render_pdf", render)
    stage = (60.0, 190.0, 1160.0, 460.0)
    header = (60.0, 38.0, 1160.0, 60.0)
    with pptx_style.render_artwork(
        pptx_import.normalize_pptx(style_source()), [0, 1], stage, header
    ) as art:
        assert len(art) == 2
        assert not art[1].get_text().strip()
    cleaned = Presentation(BytesIO(captured[0]))
    names = [s.name for s in cleaned.slides[1].shapes]
    assert names == ["edge bar"]
    original = Presentation(BytesIO(style_source()))
    assert "content photo" in [s.name for s in original.slides[1].shapes]


def test_accent_comes_from_colors_the_deck_uses():
    prs = Presentation(BytesIO(style_source()))
    measured = dict(
        title=dict(color="#182F36"),
        palette=Counter({"#182F36": 40, "#00F2FF": 6, "#7A7A7A": 30}),
        card=None,
    )
    # A dark title is not an accent; the vivid text color used in pages is.
    assert pptx_style.pick_accent(prs, measured) == "#00F2FF"
    measured["title"]["color"] = "#F59E0B"
    assert pptx_style.pick_accent(prs, measured) == "#F59E0B"


def test_readable_keeps_source_colors_and_replaces_illegible_ones():
    assert pptx_style.readable("#94A3B8", "#141619") == "#94A3B8"
    assert pptx_style.readable("#2D1410", "#2D1410") == "#FFFFFF"
    assert pptx_style.readable("#F1F5F9", "#FFFFFF") == "#111111"


def test_cards_are_sized_by_their_budget_instead_of_stretched(monkeypatch):
    monkeypatch.setattr(pptx_import, "render_pdf", fake_render)
    package, _ = pptx_style.style_replicate_pptx(style_source(), "Deck style")
    component = next(c for c in package.components if c.id == "points_3")
    root = etree.fromstring(component.svg.encode())
    cards = [n for n in root.iter("{*}rect") if n.get("rx")]
    assert len(cards) == 3
    width = float(cards[0].get("width"))
    expected = pptx_style.point_height(width, 90)
    stage_height = pptx_style.DEFAULT_STAGE[3]
    assert all(abs(float(c.get("height")) - expected) < 1 for c in cards)
    assert expected < stage_height
    limits = slot_limits(component)
    # The body box still holds its whole budget, so generated copy fills the card.
    assert all(
        limits[f"blocks[{i}].body"] == 90 for i in range(3)
    ), limits
